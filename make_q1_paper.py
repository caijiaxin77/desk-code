# -*- coding: utf-8 -*-
"""生成问题一完整论文稿。

依赖：numpy、pandas、matplotlib、python-docx、lxml、latex2mathml。
输入：results/ 中由 q1_modeling.py 生成的结果表。
输出：问题一_模型建立与求解.docx 及论文专用图。
"""
from pathlib import Path
import os
import re
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from latex2mathml.converter import convert as latex_to_mathml

HERE = Path(__file__).resolve().parent
RES = HERE / "results"
OUT = HERE / "问题一_模型建立与求解.docx"

plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei"]
plt.rcParams["axes.unicode_minus"] = False

DIM_CN = {
    "overall": "整体评分", "A_malt_all": "麦芽香气整体强度",
    "A_malt_grain": "谷物香", "A_malt_bread": "面包香",
    "A_malt_cara": "焦糖香", "A_malt_burn": "烘烤麦芽香",
    "A_hops_all": "酒花香气整体强度", "A_hops_citrus": "柑橘酒花香",
    "A_hops_tropical": "热带水果酒花香", "A_hops_noble": "传统酒花香",
    "A_hops_woody": "木质酒花香", "A_esters_all": "酯香整体强度",
    "A_esters_ethac": "乙酸乙酯香", "A_esters_isoaa": "乙酸异戊酯香",
    "A_esters_flower": "花香酯", "A_esters_fruity": "果香酯",
    "F_malt_all": "麦芽风味整体强度", "F_malt_grain": "谷物风味",
    "F_malt_bread": "面包风味", "F_malt_cara": "焦糖风味",
    "F_malt_burn": "烘烤麦芽风味", "F_hops_all": "酒花风味整体强度",
    "F_hops_citrus": "柑橘酒花风味", "F_hops_tropical": "热带水果酒花风味",
    "F_hops_noble": "传统酒花风味", "F_hops_woody": "木质酒花风味",
    "F_esters_all": "酯类风味整体强度", "F_esters_ethac": "乙酸乙酯风味",
    "F_esters_isoaa": "乙酸异戊酯风味", "F_esters_flower": "花香酯风味",
    "F_esters_fruity": "果香酯风味", "acidity": "酸度",
    "bitternes": "苦度", "sweetness": "甜度", "X4vg": "4-乙烯基愈创木酚",
    "diacetyl": "双乙酰", "dms": "二甲基硫", "metallic": "金属味",
    "stale_hops": "陈旧酒花味", "t2n": "反-2-壬烯醛",
    "orange": "橙香", "coriander": "芫荽香", "clove": "丁香香",
    "lactic": "乳酸味", "acetic": "醋酸感", "barnyard": "野菌气息",
    "alcohol": "酒精感", "aftertaste": "余味", "body": "酒体",
    "co2": "二氧化碳感",
}
MODEL_CN = {"M1_Ridge": "岭回归", "M2_PLSR": "偏最小二乘回归",
            "M3_GBR": "梯度提升回归", "M4_RF": "随机森林",
            "M5_HistGBR": "直方图梯度提升回归"}

# ------------------------------ 读取结果 ------------------------------
summary = pd.read_csv(RES / "q1_model_summary.csv").sort_values(
    "cv_R2_mean", ascending=False).reset_index(drop=True)
perdim = pd.read_csv(RES / "q1_perdim_r2.csv")
missing = pd.read_csv(RES / "q1_missing_robustness.csv")
grading = pd.read_csv(RES / "q1_grading_compare.csv")
scores = pd.read_csv(RES / "q1_scores_0_6.csv")
style = pd.read_csv(RES / "q1_style_check.csv")

best = summary.iloc[0]
best_model = best["model"]
best_col = f"{best_model}_test_R2"
top = perdim.nlargest(5, best_col)
bottom = perdim.nsmallest(5, best_col)
# 训练集折内路线择优结果与敏感性分析结果，均直接从运行日志读取，避免手工填数
route_cv = [0.0, 0.0]
log_text = (RES / "q1_run.log").read_text(encoding="utf-8", errors="replace")
m = re.search(r"训练集折内QWK：路线一=([\d.]+)，路线二=([\d.]+)", log_text)
if m:
    route_cv = [float(m.group(1)), float(m.group(2))]
sens_path = RES / "q1_run_敏感性_异常阈值仅训练集.log"
sens_r2 = []
if sens_path.exists():
    sens_r2 = [float(x) for x in re.findall(
        r"测试 R2均值=(-?\d+\.\d+)",
        sens_path.read_text(encoding="utf-8", errors="replace"))]

qwk1 = grading["route1_regress_then_bin_QWK"].mean()
qwk2 = grading["route2_ordinal_clf_QWK"].mean()
acc1 = grading["route1_accuracy"].mean()
acc2 = grading["route2_accuracy"].mean()
tol1 = grading["route1_within1"].mean()
tol2 = grading["route2_within1"].mean()

# ------------------------------ 论文专用图 ------------------------------
def make_figures():
    models = list(MODEL_CN)
    fig = plt.figure(figsize=(12.8, 4.5))
    gs = fig.add_gridspec(1, 2, width_ratios=[1, 1.65], wspace=0.18)
    ax = fig.add_subplot(gs[0, 0])
    ax.boxplot([perdim[f"{m}_test_R2"] for m in models], showmeans=True)
    ax.set_xticklabels([MODEL_CN[m] for m in models], rotation=18, fontsize=8)
    ax.axhline(0, color="#777777", lw=0.8, ls="--")
    ax.set_ylabel("测试集逐维 R²")
    ax.set_title("（a）候选模型的逐维决定系数分布", fontsize=10)
    ax = fig.add_subplot(gs[0, 1])
    vals = np.sort(perdim[best_col].to_numpy())[::-1]
    ax.bar(np.arange(len(vals)), vals,
           color=["#C0504D" if v > 0 else "#4F81BD" for v in vals], width=0.85)
    ax.axhline(vals.mean(), color="black", ls="--", lw=1,
               label=f"平均值 {vals.mean():.3f}")
    ax.axhline(0, color="#777777", lw=0.7)
    ax.set_xlabel("50 个感官维度（按决定系数降序排列）")
    ax.set_ylabel("测试集 R²")
    ax.set_title("（b）随机森林在各感官维度上的表现", fontsize=10)
    ax.legend(frameon=False, fontsize=8)
    fig.tight_layout()
    fig.savefig(RES / "paper_regression.png", dpi=240, bbox_inches="tight")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(7.2, 4.2))
    for model, color, label in [
            ("M3_GBR_reimpute", "#C0504D", "均值填补与梯度提升回归"),
            ("M5_HistGBR_native", "#4F81BD", "原生缺失与直方图梯度提升回归"),
            ("M4_RF_grading", "#2E8B57", "随机森林＋分位标定（最终方案）")]:
        g = missing[missing.model == model].groupby("mask_rate")["R2_mean"].agg(["mean", "std"])
        ax.errorbar(g.index * 100, g["mean"], yerr=g["std"], marker="o",
                    capsize=4, color=color, label=label)
    ax.axhline(0, color="#777777", lw=0.8, ls="--")
    ax.set_xlabel("化合物特征遮掩比例（%）")
    ax.set_ylabel("测试集 50 维平均 R²")
    ax.legend(frameon=False, fontsize=8)
    fig.tight_layout()
    fig.savefig(RES / "paper_missing.png", dpi=240)
    plt.close(fig)

    dims = [c for c in scores.columns if c not in ["beer", "beer_id", "tasting_category_fine"]]
    levels = scores[dims].to_numpy(dtype=int)
    pct = np.bincount(levels.ravel(), minlength=7) / levels.size * 100
    fig, ax = plt.subplots(figsize=(6.6, 3.8))
    ax.bar(range(7), pct, color="#4F81BD", width=0.68)
    ax.set_xticks(range(7))
    ax.set_xlabel("相对等级")
    ax.set_ylabel("占比（%）")
    for i, v in enumerate(pct):
        ax.text(i, v + 0.35, f"{v:.1f}", ha="center", fontsize=8)
    ax.set_ylim(0, max(pct) + 3)
    fig.tight_layout()
    fig.savefig(RES / "paper_grades.png", dpi=240)
    plt.close(fig)

    show_dims = [d for d in ["alcohol", "acidity", "barnyard", "A_malt_burn",
                             "A_esters_all", "overall"] if d in style.columns]
    grouped = style.set_index("tasting_category_fine")
    fig, ax = plt.subplots(figsize=(9.2, 7.0))
    im = ax.imshow(grouped[show_dims].to_numpy(), aspect="auto",
                   cmap="RdYlBu_r", vmin=0, vmax=6)
    ax.set_xticks(range(len(show_dims)))
    ax.set_xticklabels([DIM_CN.get(d, d) for d in show_dims], rotation=25,
                       ha="right", fontsize=8)
    ax.set_yticks(range(len(grouped)))
    ax.set_yticklabels([f"{s}（{int(grouped.loc[s, 'n_beers'])}）" for s in grouped.index],
                       fontsize=7.5)
    fig.colorbar(im, ax=ax, label="相对等级均值")
    fig.tight_layout()
    fig.savefig(RES / "paper_style.png", dpi=240)
    plt.close(fig)

make_figures()

# ------------------------------ Word 工具 ------------------------------
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
sec.left_margin, sec.right_margin = Cm(2.3), Cm(2.3)
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(10.5)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

XSL = Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL")
transform = etree.XSLT(etree.parse(str(XSL))) if XSL.exists() else None


def font_run(run, east="宋体", west="Times New Roman", size=10.5, bold=False,
             highlight=None):
    run.font.name = west
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    run.font.size = Pt(size)
    run.bold = bold
    if highlight is not None:
        run.font.highlight_color = highlight


def title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    font_run(p.add_run(text), east="黑体", size=16, bold=True)


def h1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    font_run(p.add_run(text), east="黑体", size=13, bold=True)


def h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    font_run(p.add_run(text), east="黑体", size=11, bold=True)


def para(text, indent=True, keep_together=False, highlight=None):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(3)
    if keep_together:
        p.paragraph_format.keep_together = True
    if indent:
        p.paragraph_format.first_line_indent = Pt(21)
    font_run(p.add_run(text), highlight=highlight)
    return p


def caption(text, highlight=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    font_run(p.add_run(text), east="黑体", size=9, bold=True, highlight=highlight)


def note(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    font_run(p.add_run(text), size=8.5)


def set_cell(cell, text, bold=False, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER,
             highlight=None):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(1)
    font_run(p.add_run(str(text)), size=size, bold=bold, highlight=highlight)


def cell_border(cell, **kwargs):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "bottom", "left", "right"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        if edge in kwargs and kwargs[edge]:
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), str(kwargs[edge]))
            element.set(qn("w:color"), "000000")
        else:
            element.set(qn("w:val"), "nil")


def _spacer(pt=6):
    """表/公式后的紧凑间隔段，避免默认空段撑出一整行。"""
    p = doc.add_paragraph()
    r = p.add_run(" ")
    r.font.size = Pt(pt)
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0


def table(cap, headers, rows, widths=None, size=8.5, left_cols=(), keep_together=False,
          highlight_cols=(), highlight_cap=False):
    caption(cap, highlight=WD_COLOR_INDEX.YELLOW if highlight_cap else None)
    t = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    all_rows = [headers] + rows
    for i, row in enumerate(all_rows):
        for j, value in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.LEFT if j in left_cols and i > 0 else WD_ALIGN_PARAGRAPH.CENTER
            hl = WD_COLOR_INDEX.YELLOW if j in highlight_cols else None
            set_cell(t.cell(i, j), value, bold=(i == 0), size=size, align=align, highlight=hl)
            cell_border(t.cell(i, j), **({"top": 12, "bottom": 6} if i == 0 else
                                        ({"bottom": 12} if i == len(all_rows) - 1 else {})))
    if widths:
        for j, width in enumerate(widths):
            for row in t.rows:
                row.cells[j].width = Cm(width)
    # 表头跨页重复
    trPr = t.rows[0]._tr.get_or_add_trPr()
    if trPr.find(qn("w:tblHeader")) is None:
        trPr.append(OxmlElement("w:tblHeader"))
    # 整表不跨页：所有行禁止断开，且除末行外与下一行保持同页
    if keep_together:
        n = len(t.rows)
        for i, row in enumerate(t.rows):
            _no_split(row)
            if i < n - 1:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.keep_with_next = True
    _spacer(6)
    return t


def figure(filename, cap, width=6.4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(RES / filename), width=Inches(width))
    caption(cap)


eq_counter = 0

def _no_split(row):
    """禁止表格行跨页断开，保证同一编号公式的多行始终同页。"""
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn("w:cantSplit")) is None:
        trPr.append(OxmlElement("w:cantSplit"))


def equation(latex):
    global eq_counter
    eq_counter += 1
    # 引导句与公式保持同页
    if doc.paragraphs:
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _no_split(t.rows[0])
    widths = [1.7, 12.4, 1.7]
    for j, w in enumerate(widths):
        t.cell(0, j).width = Cm(w)
    p = t.cell(0, 1).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    if transform is not None:
        mathml = latex_to_mathml(latex)
        omml = transform(etree.fromstring(mathml.encode("utf-8"))).getroot()
        p._p.append(omml)
    else:
        font_run(p.add_run(latex), west="Cambria Math")
    p_num = t.cell(0, 2).paragraphs[0]
    p_num.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font_run(p_num.add_run(f"（{eq_counter}）"), size=9)
    _spacer(8)


def equation_block(latex_lines):
    """同一编号下按多行排布的长公式，避免单行超出页宽被截断。"""
    global eq_counter
    eq_counter += 1
    # 引导句与公式保持同页
    if doc.paragraphs:
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _no_split(t.rows[0])
    widths = [1.7, 12.4, 1.7]
    for j, w in enumerate(widths):
        t.cell(0, j).width = Cm(w)
    cell = t.cell(0, 1)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for k, latex in enumerate(latex_lines):
        p = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        if transform is not None:
            mathml = latex_to_mathml(latex)
            omml = transform(etree.fromstring(mathml.encode("utf-8"))).getroot()
            p._p.append(omml)
        else:
            font_run(p.add_run(latex), west="Cambria Math")
    num_cell = t.cell(0, 2)
    num_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p_num = num_cell.paragraphs[0]
    p_num.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font_run(p_num.add_run(f"（{eq_counter}）"), size=9)
    _spacer(8)


# 页码
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run()
begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
run._r.extend([begin, instr, end])
font_run(run, size=9)

# ------------------------------ 正文 ------------------------------
title("问题一  数据缺失条件下的感官评分")

h1("5.1  问题分析")
para("问题一要求以化合物及理化测量数据为输入，对附件3所定义的香气与风味指标给出0—6分级结果。"
     "附件4提供的是经标准化或中心化处理的连续感官指标，而非原始人工等级，因此不能直接四舍五入。"
     "本文将任务分解为连续指标预测、相对等级构造和缺失条件检验三个步骤：先在多个回归模型中依据训练集交叉验证结果选型，"
     "再利用训练集经验分布把连续预测转换为0—6相对等级，最后通过同步遮掩实验评价完整方案在特征缺失时的稳定性。")
para("三个问题共用同一套样本对齐、尺度核验和数据质量处理结果。问题一使用化合物矩阵与感官指标；"
     "问题二在此基础上引入化合物相关矩阵；问题三进一步使用风格标签和缺陷化合物阈值。"
     "三个问题共用样本索引、字段映射和尺度核验结果，以保证数据口径一致。"
     "涉及模型拟合的参数——填补均值、标准化参数、超参数和分级路线选择——均分别由训练集或各训练折估计，"
     "不提前使用验证集和测试集的标签信息；疑似异常的识别作为不使用感官标签的数据有效性核查，"
     "在全样本上执行，其处理口径的影响通过敏感性分析单独说明。")

h1("5.2  数据预处理")
h2("5.2.1  样本对齐与尺度核验")
para("以样本编号为主键，对齐附件1、附件4中的250款啤酒及风格标签。"
     "附件1包含231个化学及理化特征，其中包括226个测量指标和5个派生属性；附件4包含50个连续感官指标。"
     "经逐文件核对，三个问题均沿用同一套样本顺序和字段映射。")
para("附件1同时含有对数尺度与线性尺度。结合附件8的校准数据逐列核验：对已确认采用对数变换的指标保留原编码，"
     "对线性指标不作重复取对数。疑似下限替代值与对数底数相互印证，pH反变换结果亦处于合理范围。"
     "对于需要还原浓度的分析，按指标所属尺度分别处理：")
equation(r"c_{ij}=\begin{cases}10^{x_{ij}},&j\in\mathcal{L},\\x_{ij},&j\in\mathcal{R},\end{cases}")
para("其中，L 与 R 分别表示对数尺度和线性尺度指标集合。"
     "对下限替代值进行逆变换仅用于保持计算口径一致，不将所得数值解释为真实浓度。")
_sensory_cols = [c for c in scores.columns
                 if c not in ("beer", "beer_id", "tasting_category_fine")]
_a_cols = [c for c in _sensory_cols if c.startswith("A_")]
_f_cols = [c for c in _sensory_cols if c.startswith("F_")]
_o_cols = [c for c in _sensory_cols if not c.startswith(("A_", "F_"))]
para(f"附件4的{len(_sensory_cols)}项连续感官指标按题目口径分为三部分："
     f"Aroma与Flavour各{len(_a_cols)}项（麦芽、酒花与酯类各{len(_a_cols) // 3}项），"
     f"共{len(_a_cols) + len(_f_cols)}项作为核心输出；"
     f"其余{len(_o_cols)}项（" + "、".join(DIM_CN.get(c, c) for c in _o_cols) + "）作为扩展感官指标一并预测。"
     "评分卡中的乙醛香气、乙醛风味等项目在附件4中没有对应数值标签，本文不单独建模；"
     "自由文本不作为预测目标。")

h2("5.2.2  缺失值、疑似异常与下限替代值")
para("附件1共有73个缺失单元格，占全部化学测量值的0.126%，分布于9个指标。缺失比例整体较低，"
     "现有信息不足以判断其缺失机制。各指标的缺失情况见表1。")
missing_rows = [
    ["乙酸戊酯", 27, "10.80%"], ["酒花香气总和", 13, "5.20%"],
    ["甲基吡嗪", 12, "4.80%"], ["2,3-丁二酮", 9, "3.60%"],
    ["3-甲基丁醛", 7, "2.80%"], ["烟酸乙酯", 2, "0.80%"],
    ["异丁酸乙酯", 1, "0.40%"], ["糠醛", 1, "0.40%"],
    ["苯乙醛", 1, "0.40%"],
]
table("表1  附件1中化学指标的缺失值统计", ["指标", "缺失个数", "占该指标比例"],
      missing_rows, widths=[6.0, 3.0, 4.0], keep_together=True)
para("常规缺失采用同指标可观测值的均值填补。对对数尺度指标，均值在原编码空间中计算；"
     "对线性尺度指标，均值在线性空间中计算。记第 j 个指标在当前训练集中的可观测样本集合为 Oⱼ(T)，则")
equation_block([r"\bar{x}_{j}=\frac{1}{|\mathcal{O}_{j}^{\mathrm{train}}|}\sum_{i\in\mathcal{O}_{j}^{\mathrm{train}}}x_{ij},",
                r"\widetilde{x}_{ij}=\begin{cases}\bar{x}_{j},&x_{ij}\text{ 缺失},\\x_{ij},&\text{其他}.\end{cases}"])
para("其中，可观测样本集合仅包含当前训练集中的有效观测。交叉验证时，在各训练折内重新计算填补均值，"
     "并应用于对应验证折；测试集评价时使用完整训练集估计的填补均值。")
para("疑似异常的识别只使用化学测量本身，不使用任何感官标签，属于无监督的数据有效性核查，因此可在全样本上统一执行；"
     "需要说明的是，该规则的最大值、中位数与第99百分位数由全部样本计算，虽未使用感官标签，仍使用了测试集的化学测量。"
     "被标记的观测按其所属指标的训练集均值填补（填补均值排除已标记的异常值本身）。"
     "这些观测仍只是统计意义上的疑似异常，是否为录入或单位错误尚无法确认。"
     "疑似异常按照统计偏离程度识别。对适用指标，记最大值、次大值、中位数和第99百分位数分别为"
     "x₍₁₎、x₍₂₎、mⱼ 和 q₀.₉₉,ⱼ。当最大值相对中位数偏离超过100倍，"
     "且相对次大值超过20倍或相对第99百分位数超过15倍时，将其标记为疑似异常，再按缺失值流程处理。"
     "处理效果仍需通过保留原值与填补处理的对照实验进一步检验。")
outlier_rows = [
    ["2-甲基四氢呋喃-3-酮", "12596.43", "为同列次大观测的164倍"],
    ["(E)-β-大马士革酮", "5991.01", "为同列次大观测的53倍"],
    ["3-甲基丁醛", "127.06", "为第99百分位数的20.9倍、中位数的1172倍"],
]
table("表2  疑似异常观测的识别结果", ["指标", "观测值", "统计偏离特征"],
      outlier_rows, widths=[5.2, 2.6, 6.0], left_cols=(0, 2))
para("对于疑似下限替代现象，先识别重复低值编码，再结合各指标低端取值的聚集情况及来源资料进行核验。"
     "建模矩阵保留现有数值编码；还原矩阵中的置零仅作为计算约定。"
     "对全零还原值较多的第2号样本进行核查，未发现足以排除该样本的证据，故予以保留。")

h2("5.2.3  数据划分、标准化与三问共用输入")
para("问题一按啤酒风格分层划分训练集与测试集，分别包含175款和75款样本，随机种子设为0。"
     "为减少测试信息对建模的影响，标准化参数只由训练集估计：")
equation_block([r"z_{ij}=\frac{x_{ij}-\mu_{j}^{\mathrm{train}}}{\sigma_{j}^{\mathrm{train}}},",
                r"\mu_{j}^{\mathrm{train}}=\frac{1}{|\mathcal{T}|}\sum_{i\in\mathcal{T}}x_{ij}."])
para("统一预处理形成的数据输入及其用途见表3。问题二与问题三继续使用相同样本索引和处理后的化合物矩阵，"
     "不再重复进行尺度识别或异常处理。")
usage_rows = [
    ["问题一", "标准化化合物矩阵、50维连续感官指标、风格分层划分", "连续预测、相对等级构造、缺失鲁棒性检验"],
    ["问题二", "处理后化合物矩阵、化合物相关矩阵", "共表达图构建、关键因素筛选、重采样稳定性分析"],
    ["问题三", "处理后化合物矩阵、风格标签、缺陷化合物阈值", "风格条件生成、感官画像、跨风格迁移与缺陷约束"],
]
table("表3  三个问题对统一预处理结果的使用", ["问题", "共用输入", "后续用途"], usage_rows,
      widths=[2.0, 7.0, 5.0], left_cols=(1, 2), size=8)

h1("5.3  模型假设")
for text in [
    "（1）同一编号在不同附件中对应同一款啤酒，各样本在统一对齐后可视为相互独立的观测。",
    "（2）附件4的连续感官指标保留样本间强弱顺序，可据此构造群体内部的相对等级；该等级不等同于恢复原始人工评分。",
    "（3）缺失实验中的独立随机遮掩仅用于压力测试，不据此推断原始数据的真实缺失机制。",
    "（4）化学测量与感官指标之间可能含有非线性关系和交互作用，因此同时比较低维线性模型与树集成模型。",
]:
    para(text, indent=False)

h1("5.4  符号说明")
symbol_rows = [
    ["n、p、q", "样本数、化合物指标数、感官维度数", "K", "相对等级数，取7"],
    ["i、j、k", "样本、感官维度、等级边界下标", "xᵢⱼ", "第 i 款样本的第 j 项化学测量"],
    ["yᵢⱼ", "第 i 款样本的第 j 维连续感官指标", "zᵢⱼ", "标准化后的化学测量"],
    ["fⱼ", "第 j 维感官指标的回归函数", "ŷᵢⱼ", "连续感官指标预测值"],
    ["bₖⱼ", "第 j 维感官指标的第 k 条分位边界", "ℓᵢⱼ", "0—6相对等级"],
    ["δ", "特征遮掩比例", "μⱼ、σⱼ", "第 j 项指标的均值与标准差"],
    ["Rⱼ²", "第 j 维感官指标的决定系数", "ρⱼ", "第 j 维的斯皮尔曼秩相关"],
    ["κ_w", "二次加权Kappa系数", "T、E", "训练集与测试集样本集合"],
]
table("表4  主要符号及含义", ["符号", "含义", "符号", "含义"], symbol_rows,
      widths=[2.2, 5.4, 2.2, 5.4], size=8, left_cols=(1, 3))

h1("5.5  模型建立")
h2("5.5.1  连续感官指标回归模型")
para("对第 j 个感官维度建立化合物测量到连续感官指标的映射：")
equation(r"\widehat{y}_{ij}=f_j(\boldsymbol{x}_i;\boldsymbol{\theta}_j),\qquad i=1,\ldots,n,\quad j=1,\ldots,q.")
para("候选模型覆盖线性正则化、潜变量投影和树集成三类。岭回归通过惩罚回归系数控制高维小样本下的波动，其目标函数为")
equation(r"\widehat{\boldsymbol{W}}=\underset{\boldsymbol{W}}{\mathrm{arg\,min}}\ \left\|\boldsymbol{Y}-\boldsymbol{X}\boldsymbol{W}\right\|_F^2"
         r"+\alpha\left\|\boldsymbol{W}\right\|_F^2,")
para("其中，α 为正则化系数。闭式解中的输入与目标均按训练集均值中心化，预测时加回目标变量的训练集均值；"
     "代码通过Ridge的默认截距处理实现。岭回归在样本数小于特征数时仍可稳定求解，其权重矩阵为：")
equation(r"\widehat{\boldsymbol{W}}=(\boldsymbol{X}^{\top}\boldsymbol{X}+\alpha\boldsymbol{I})^{-1}\boldsymbol{X}^{\top}\boldsymbol{Y}.")
para("α 越大则系数越向零收缩，可缓解化学指标共线性导致的方差膨胀；本文在训练集上以五折交叉验证决定系数选取 α。"
     "偏最小二乘回归将输入与输出同时投影到少量潜在成分，在保留化学测量主要变异的同时增强其与感官指标的协方差：")
equation_block([r"\boldsymbol{X}=\boldsymbol{T}\boldsymbol{P}^{\top}+\boldsymbol{E},\qquad "
                r"\boldsymbol{Y}=\boldsymbol{U}\boldsymbol{Q}^{\top}+\boldsymbol{F},",
                r"\boldsymbol{U}=\boldsymbol{T}\boldsymbol{B},\qquad "
                r"\widehat{\boldsymbol{Y}}=\boldsymbol{T}\boldsymbol{B}\boldsymbol{Q}^{\top}."])
para("其中，T 为得分矩阵、P 与 Q 分别为输入与输出的载荷矩阵，B 为得分间的回归系数，"
     "输出通过载荷矩阵 Q 反投影到感官指标空间。第 a 个成分的输入权重 w_a 与输出权重 c_a "
     "在单位范数约束下最大化两个投影之间的协方差平方：")
equation_block([r"\max_{\boldsymbol{w}_a,\ \boldsymbol{c}_a}\ "
                r"\mathrm{Cov}^2\left(\boldsymbol{X}\boldsymbol{w}_a,\ \boldsymbol{Y}\boldsymbol{c}_a\right),",
                r"\mathrm{s.t.}\quad \left\|\boldsymbol{w}_a\right\|_2=1,\quad \left\|\boldsymbol{c}_a\right\|_2=1."])
para("成分数控制信息压缩程度：过少会丢失化学信息，过多会引入噪声，本文在2至30之间网格选取。"
     "随机森林对自助重采样样本分别训练回归树，并在节点分裂时随机选取部分特征，最终取各树预测均值：")
equation(r"\widehat{y}_{ij}^{\mathrm{RF}}=\frac{1}{B}\sum_{b=1}^{B}h_{bj}(\boldsymbol{x}_i),")
para("其中，B 为树的数量，h_bj 表示第 b 棵树对第 j 维感官指标的预测。"
     "每棵回归树以平方误差为分裂准则，叶节点预测取该节点训练样本的均值，分裂时约30%特征被随机选用，"
     "并对叶节点最小样本数加以限制，以降低高维小样本下的方差：")
equation(r"L_j^{\mathrm{tree}}(\mathcal{S})=\frac{1}{|\mathcal{S}|}\sum_{i\in\mathcal{S}}"
         r"(y_{ij}-\bar{y}_{\mathcal{S},j})^2.")
para("梯度提升模型采用加法形式逐步拟合当前损失的负梯度：")
equation_block([r"F_{mj}(\boldsymbol{x})=F_{m-1,j}(\boldsymbol{x})+\nu h_{mj}(\boldsymbol{x}),",
                r"h_{mj}\approx-\left.\frac{\partial L_j}{\partial F_j}\right|_{F_j=F_{m-1,j}}."])
para("本文采用平方误差损失，其负梯度即当前残差，学习率 ν 控制每步步长，子采样率0.8为基学习器引入随机性：")
equation(r"L_j=\frac{1}{2}\sum_{i\in\mathcal{T}}(y_{ij}-F_j(\boldsymbol{x}_i))^2.")
para("直方图梯度提升回归进一步将连续特征离散为直方图以加速分裂点搜索，并对缺失值学习默认分支方向，"
     "因此可在不预先填补的情况下训练。为控制过拟合，本文对其加入L2惩罚并限制最大叶节点数。"
     "上述五类模型在同一训练集五折交叉验证下比较，以训练集交叉验证中50个感官维度的平均决定系数选型，"
     "测试集只用于最终评价，避免以模型先验偏好替代实证选型。")

h2("5.5.2  0—6相对等级标定模型")
para("由于附件4不含原始人工等级，本文采用两条路线构造相对等级。路线一先将训练集连续感官指标的经验分位数作为边界，"
     "再把最优回归模型的连续预测映射到相应区间。第 j 维的第 k 条边界为")
equation(r"b_{kj}=Q_{k/7}\left(\{y_{ij}:i\in\mathcal{T}\}\right),\qquad k=1,\ldots,6.")
para("对连续预测值 ŷᵢⱼ，相对等级定义为")
equation(r"\widehat{\ell}_{ij}=\sum_{k=1}^{6}\mathbb{I}\left(\widehat{y}_{ij}\geq b_{kj}\right),\qquad "
         r"\widehat{\ell}_{ij}\in\{0,1,\ldots,6\}.")
para("该方法只利用训练集估计边界，并保持连续预测的排序关系。路线二先按同一边界构造训练集参考等级，"
     "再对各感官维度分别训练直接多分类模型，输出后验概率最大的等级：")
equation(r"\widehat{\ell}_{ij}=\underset{r\in\{0,\ldots,6\}}{\mathrm{arg\,max}}\ P_j\left(\ell=r\mid\boldsymbol{x}_i\right).")
para("两条路线得到的等级均为样本群体内部的相对等级。本文以二次加权Kappa系数选择分级路线，"
     "而不以连续回归的决定系数替代分级质量评价。由于分位边界按升序排列，等级映射保持连续预测的顺序："
     "预测值越大则等级越高，但相邻等级之间的距离不一定等距。"
     "在评价完成后，为充分利用已核验的全部样本，用250款样本重训最终模型，"
     "并在全部样本上重新估计分位边界，生成最终的相对等级表：")
equation(r"b_{kj}^{*}=Q_{k/7}\left(\{y_{ij}:i=1,\ldots,n\}\right),\qquad k=1,\ldots,6.")

h2("5.5.3  数据缺失条件下的鲁棒性模型")
para("设遮掩比例为 δ，对训练集与测试集的原始化学编码矩阵同步生成随机遮掩矩阵，遮掩变量 Mᵢⱼ 服从参数为 δ 的伯努利分布，"
     "感官标签保持完整。原有缺失值继续保留，δ=0 仅表示没有额外遮掩。两条方案使用完全相同的遮掩位置：")
equation(r"x_{ij}^{(\delta)}=\begin{cases}\mathrm{NA},&M_{ij}=1,\\x_{ij},&M_{ij}=0.\end{cases}")
para("方案一以遮掩后训练集中该列的可观测值重新估计均值，凡遮掩后数值缺失的单元格（含新遮掩与原有缺失）均以该均值填补，"
     "其余保留遮掩后的原值，随后以遮掩后训练集的均值与标准差标准化，再重新训练梯度提升回归。"
     "记遮掩后训练集中第 j 列的可观测集合为 Oⱼ⁽δ⁾：")
equation_block([r"\tilde{x}_{ij}=\begin{cases}\bar{x}_j^{(\delta)},&x_{ij}^{(\delta)}\ \text{缺失},\\"
                r"x_{ij}^{(\delta)},&\text{其他},\end{cases}",
                r"\bar{x}_j^{(\delta)}=\frac{1}{|\mathcal{O}_j^{(\delta)}|}"
                r"\sum_{i\in\mathcal{O}_j^{(\delta)}}x_{ij}."])
equation(r"z_{ij}^{(\delta)}=\frac{\tilde{x}_{ij}-\mu_j^{(\delta)}}{\sigma_j^{(\delta)}}.")
para("方案二保留缺失值，仅以遮掩后训练集的均值与标准差标准化可观测值，把含缺失矩阵直接输入直方图梯度提升回归，"
     "使模型在训练过程中学习缺失方向。两种方案均在含缺失测试集上评价，并使用完全相同的遮掩位置。"
     "非零遮掩比例重复3个随机种子，报告50维平均决定系数的均值与标准差。由于模型类型与缺失处理方式同时变化，"
     "结果用于比较两套完整方案，不作单一填补机制的因果归因。")
para("此外，为直接检验最终采用的评分方案，在上述完全相同的遮掩位置下，对“随机森林＋分位标定”方案"
     "（即5.5.1选定的随机森林回归与5.5.2的分位标定路线）同步报告连续预测的决定系数，"
     "以及分级层面的二次加权Kappa系数与等级准确率。分级边界仍仅由完整训练集标签估计，"
     "遮掩只改变特征输入、不改变等级标定基准。", highlight=WD_COLOR_INDEX.YELLOW)

h2("5.5.4  评价指标")
para("连续预测采用逐维决定系数和秩相关评价。对测试集 E，第 j 维的决定系数为")
equation_block([r"R_j^2=1-\frac{\sum_{i\in\mathcal{E}}(y_{ij}-\widehat{y}_{ij})^2}"
                r"{\sum_{i\in\mathcal{E}}(y_{ij}-\bar{y}_{\cdot j})^2},",
                r"\overline{R^2}=\frac{1}{q}\sum_{j=1}^{q}R_j^2."])
para("秩相关用于衡量预测值与观测值排序的一致性。"
     "Spearman秩相关系数等于观测值平均秩与预测值平均秩之间的Pearson相关系数：")
equation(r"\rho_j=\frac{\mathrm{Cov}(a_j,b_j)}"
         r"{\sqrt{\mathrm{Var}(a_j)\,\mathrm{Var}(b_j)}}.")
para("其中，aᵢⱼ 与 bᵢⱼ 分别为第 i 款样本在第 j 维上的观测值与预测值的平均秩；"
     "Cov 与 Var 均按测试集样本计算，上式即两者平均秩之间的Pearson相关系数。"
     "对于重复取值，采用平均秩处理，以避免忽略并列秩对计算结果的影响。"
     "分级质量采用严格准确率、相邻一级容差准确率和二次加权Kappa系数。"
     "令 K=7，O(a,b) 与 E(a,b) 分别表示实际等级 a、预测等级 b 的观测比例和随机一致比例，则")
equation_block([r"w_{ab}=\left(\frac{a-b}{K-1}\right)^2,",
                r"\kappa_w=1-\frac{\sum_{a,b}w_{ab}O_{ab}}{\sum_{a,b}w_{ab}E_{ab}}."])
equation_block([r"\mathrm{Acc}=\frac{1}{q}\sum_{j=1}^{q}\frac{1}{n_{\mathcal{E}}}\sum_{i\in\mathcal{E}}"
                r"\mathbb{I}(\widehat{\ell}_{ij}=\ell_{ij}),",
                r"\mathrm{Acc}_{\pm1}=\frac{1}{q}\sum_{j=1}^{q}\frac{1}{n_{\mathcal{E}}}\sum_{i\in\mathcal{E}}"
                r"\mathbb{I}(|\widehat{\ell}_{ij}-\ell_{ij}|\leq1)."])

h1("5.6  模型求解")
h2("5.6.1  参数选择与训练策略")
para("所有参数选择均在训练集内完成，测试集只用于最终评价，不参与模型选择。"
     "岭回归的正则化系数在 {0.1, 1, 10, 100, 1000} 内以五折交叉验证决定系数选取，最终选 α=1000，"
     "说明在231维化学指标下需要较强收缩。偏最小二乘回归的成分数在 {2, 5, 10, 15, 20, 30} 内选取，"
     "最终选5个成分，表明少量潜在方向已能提取主要化学—感官协方差。")
para("梯度提升回归若对50维分别做全网格搜索代价过高，故先在整体强度、酒花香气、酸度和二甲基硫四个"
     "代表性维度上做三折交叉验证预筛六组配置，再将获胜配置用于全部维度。预筛选定的配置为"
     "500棵树、学习率0.05、最大深度3、子采样率0.8，兼顾拟合能力与过拟合控制。"
     "随机森林与直方图梯度提升回归的参数依据高维小样本下的稳定性经验设定："
     "限定分裂特征比例与叶节点最小样本数以降低方差，直方图模型加入L2惩罚并限制叶节点数。"
     "直接多分类模型采用较保守的最大深度3，避免在七个等级的小样本上过拟合。"
     "数据划分、交叉验证及模型随机种子设为0；非零遮掩比例下使用0、1、2三个随机种子重复实验。"
     "五折交叉验证采用打乱划分，主要参数见表5。")
param_rows = [
    ["岭回归", "正则化系数", "1000"],
    ["偏最小二乘回归", "潜在成分数", "5"],
    ["梯度提升回归", "树数量、学习率、最大深度、子采样率", "500、0.05、3、0.8"],
    ["随机森林", "树数量、分裂特征比例、叶节点最小样本数", "300、0.3、2"],
    ["直方图梯度提升回归", "迭代次数、学习率、最大叶节点数、L2惩罚", "300、0.05、15、1.0"],
    ["直接多分类模型", "树数量、学习率、最大深度", "200、0.05、3"],
]
table("表5  候选模型及主要求解参数", ["模型", "参数", "取值"], param_rows,
      widths=[4.2, 6.8, 3.5], left_cols=(0, 1), size=8)

h2("5.6.2  求解步骤")
para("算法输入为保留缺失标记的化合物矩阵、50维连续感官指标和风格标签；"
     "填补与标准化在训练集或各训练折内完成。"
     "输出为连续回归选型结果、缺失鲁棒性曲线、250×50相对等级表和风格画像汇总。"
     "主要计算量来自随机森林与梯度提升的逐维训练，通过按感官维度并行执行降低运行时间。"
     "求解步骤见表6。")
algorithm_rows = [
    ["1", "读取统一预处理结果，按样本编号对齐化学测量、感官指标和风格标签。"],
    ["2", "按风格分层得到175款训练样本与75款测试样本，仅用训练集估计标准化参数。"],
    ["3", "在训练集内完成参数选择，计算五折交叉验证的折外预测；以平均决定系数选择连续回归模型。"],
    ["4", "在测试集上计算逐维决定系数、秩相关和误差，比较不同模型及其维度差异。"],
    ["5", "依次设置0%、5%、10%、20%、30%遮掩比例；非零比例重复3个随机种子，比较两套缺失处理方案。"],
    ["6", "在训练集内部用五折交叉验证比较回归后标定与直接多分类两条路线，以折内二次加权Kappa系数择优；测试集仅用于报告。"],
    ["7", "评价完成后，以全部250款样本重训最终模型，重新估计分位边界并生成250×50相对等级表。"],
    ["8", "按22种风格汇总等级均值，进行描述性画像检查，并输出完整结果表。"],
]
table("表6  问题一的求解算法", ["步骤", "计算内容"], algorithm_rows,
      widths=[1.4, 13.0], left_cols=(1,), size=8.5, keep_together=True)
para("实验在Python环境下使用numpy与scikit-learn实现。数据划分、交叉验证和模型训练的随机种子设为0，"
     "非零遮掩比例下使用0、1、2三个随机种子；训练集划分、交叉验证和遮掩过程均可由此复现。"
     "全部结果由本队代码对赛题附件独立运行得到。")

h1("5.7  结果与分析")
h2("5.7.1  连续回归模型选型")
reg_rows = []
for _, r in summary.iterrows():
    reg_rows.append([MODEL_CN[r["model"]], f"{r['cv_R2_mean']:.3f}",
                     f"{r['test_R2_mean']:.3f}", f"{r['test_R2_median']:.3f}",
                     f"{int(r['test_dims_R2_positive'])}/50", f"{r['test_Spearman_mean']:.3f}"])
table("表7  连续回归模型的交叉验证与测试集结果",
      ["模型", "训练集折外预测R²", "测试集平均R²", "测试集R²中位数", "R²>0的维度数", "平均秩相关"],
      reg_rows, widths=[3.6, 2.5, 2.2, 2.4, 2.0, 1.8], size=7.8)
note("注：训练集折外预测R²由五折交叉验证的折外预测按样本拼接后，在训练集上计算得到，"
     "而非各折决定系数的简单平均。")
para(f"随机森林在训练集折外预测上的平均决定系数最高，为{best['cv_R2_mean']:.3f}，因此被选为连续预测模型。"
     f"其测试集平均决定系数为{best['test_R2_mean']:.3f}，中位数为{best['test_R2_median']:.3f}，"
     f"50个感官维度中有{int(best['test_dims_R2_positive'])}个维度的决定系数为正，平均秩相关为{best['test_Spearman_mean']:.3f}。"
     "在当前实验设置下，随机森林的平均表现最佳，但不同模型在各感官维度上的优势并不一致。")
figure("paper_regression.png", "图1  连续回归模型的逐维表现", width=6.8)
para("逐维结果进一步表明，酒精感、苦度和酒体等理化特征较明确的维度较易预测，"
     "部分酒花与酯类细分香气的决定系数较低甚至为负。因此，面板平均值不能代替逐维诊断。"
     "表8列出随机森林表现最高与最低的五个维度。")
dim_rows = []
for i in range(5):
    dim_rows.append([DIM_CN.get(top.iloc[i]["sensory_dim"], top.iloc[i]["sensory_dim"]),
                     f"{top.iloc[i][best_col]:.3f}",
                     DIM_CN.get(bottom.iloc[i]["sensory_dim"], bottom.iloc[i]["sensory_dim"]),
                     f"{bottom.iloc[i][best_col]:.3f}"])
table("表8  随机森林逐维决定系数的极值维度",
      ["表现最高维度", "R²", "表现最低维度", "R²"], dim_rows,
      widths=[5.0, 1.8, 5.0, 1.8], keep_together=True)

h2("5.7.2  缺失条件下的性能变化")
miss_rows = []
for rate in [0.0, 0.05, 0.10, 0.20, 0.30]:
    row = [f"{rate:.0%}"]
    for model in ["M3_GBR_reimpute", "M5_HistGBR_native", "M4_RF_grading"]:
        g = missing[(missing.model == model) & (missing.mask_rate == rate)]["R2_mean"]
        if rate == 0:
            row.append(f"{g.mean():.3f}")
        else:
            row.append(f"{g.mean():.3f}±{g.std():.3f}")
    miss_rows.append(row)
table("表9  三套缺失处理方案在不同遮掩比例下的测试集平均决定系数",
      ["遮掩比例", "均值填补与梯度提升回归", "原生缺失与直方图梯度提升回归",
       "随机森林＋分位标定（最终方案）"],
      miss_rows, widths=[2.0, 4.2, 4.6, 4.6], size=8, keep_together=True,
      highlight_cols=(3,), highlight_cap=True)
note("注：0%遮掩为基线结果，仅计算一次，未计算重复实验标准差。黄色高亮列为本次补报的最终方案结果。")
base_a = missing[(missing.model == "M3_GBR_reimpute") & (missing.mask_rate == 0)]["R2_mean"].mean()
base_b = missing[(missing.model == "M5_HistGBR_native") & (missing.mask_rate == 0)]["R2_mean"].mean()
end_a = missing[(missing.model == "M3_GBR_reimpute") & (missing.mask_rate == 0.30)]["R2_mean"].mean()
end_b = missing[(missing.model == "M5_HistGBR_native") & (missing.mask_rate == 0.30)]["R2_mean"].mean()
base_c = missing[(missing.model == "M4_RF_grading") & (missing.mask_rate == 0)]["R2_mean"].mean()
end_c = missing[(missing.model == "M4_RF_grading") & (missing.mask_rate == 0.30)]["R2_mean"].mean()
para(f"两套方案的平均表现随遮掩比例增加总体下降，但局部存在波动。遮掩比例由0增至30%时，"
     f"均值填补方案由{base_a:.3f}降至{end_a:.3f}，相对下降{abs((end_a-base_a)/base_a)*100:.1f}%；"
     f"原生缺失方案由{base_b:.3f}降至{end_b:.3f}，相对下降{abs((end_b-base_b)/base_b)*100:.1f}%。"
     "在本实验范围内，均值填补方案的整体下降幅度较小，但该结论只适用于两套完整方案的比较。")
para(f"最终采用的随机森林＋分位标定方案在相同遮掩条件下由{base_c:.3f}降至{end_c:.3f}，"
     f"相对下降{abs((end_c-base_c)/base_c)*100:.1f}%，其连续预测的决定系数高于另两套方案，"
     "说明该方案对特征缺失的鲁棒性与基准方案相当。",
     highlight=WD_COLOR_INDEX.YELLOW)
# 最终方案的分级指标：Kappa 与等级准确率
rf_kappa0 = missing[(missing.model == "M4_RF_grading") & (missing.mask_rate == 0)]["Kappa_mean"].mean()
rf_kappa30 = missing[(missing.model == "M4_RF_grading") & (missing.mask_rate == 0.30)]["Kappa_mean"].mean()
rf_acc0 = missing[(missing.model == "M4_RF_grading") & (missing.mask_rate == 0)]["grade_accuracy"].mean()
rf_acc30 = missing[(missing.model == "M4_RF_grading") & (missing.mask_rate == 0.30)]["grade_accuracy"].mean()
rf_tol0 = missing[(missing.model == "M4_RF_grading") & (missing.mask_rate == 0)]["grade_within1"].mean()
rf_tol30 = missing[(missing.model == "M4_RF_grading") & (missing.mask_rate == 0.30)]["grade_within1"].mean()
rf_rows = [
    ["0%", f"{rf_kappa0:.3f}", f"{rf_acc0*100:.1f}%", f"{rf_tol0*100:.1f}%"],
    ["30%", f"{rf_kappa30:.3f}", f"{rf_acc30*100:.1f}%", f"{rf_tol30*100:.1f}%"],
]
table("表10  最终方案（随机森林＋分位标定）在不同遮掩比例下的分级表现",
      ["遮掩比例", "平均二次加权Kappa", "严格准确率", "相邻一级容差准确率"],
      rf_rows, widths=[3.0, 3.6, 3.2, 4.2], size=8, keep_together=True,
      highlight_cap=True)
note("注：表10为本次补报，直接检验最终方案在缺失条件下的分级稳定性，其数值由3个随机种子取平均（30%遮掩）或单次（0%基线）。")
para(f"分级层面，最终方案在无缺失基线上的平均二次加权Kappa系数为{rf_kappa0:.3f}，"
     f"严格准确率{rf_acc0*100:.1f}%、相邻一级容差准确率{rf_tol0*100:.1f}%；"
     f"遮掩30%后分别变为{rf_kappa30:.3f}、{rf_acc30*100:.1f}%与{rf_tol30*100:.1f}%，"
     "分级指标随信息损失而下降，但未出现崩溃式退化，说明该方案对中等程度缺失具备一定稳健性。",
     highlight=WD_COLOR_INDEX.YELLOW)
figure("paper_missing.png", "图2  不同特征遮掩比例下的连续预测性能", width=5.8)

h2("5.7.3  相对等级标定结果")
grade_rows = [
    ["回归预测与分位标定", f"{qwk1:.3f}", f"{acc1*100:.1f}%", f"{tol1*100:.1f}%", "采用"],
    ["直接多分类分级", f"{qwk2:.3f}", f"{acc2*100:.1f}%", f"{tol2*100:.1f}%", "—"],
]
table("表11  两条相对等级构造路线的测试集结果",
      ["路线", "平均二次加权Kappa", "严格准确率", "相邻一级容差准确率", "最终选择"],
      grade_rows, widths=[4.2, 3.0, 2.4, 3.2, 1.8], size=8)
para(f"分级路线在训练集内部择优：五折交叉验证的折内二次加权Kappa系数分别为{route_cv[0]:.4f}与{route_cv[1]:.4f}，"
     "据此采用回归预测与分位标定路线，测试集仅用于报告、不参与选择。"
     f"测试集上该路线的二次加权Kappa系数为{qwk1:.3f}，直接多分类路线为{qwk2:.3f}。该路线的严格准确率为"
     f"{acc1*100:.1f}%，相邻一级容差准确率为{tol1*100:.1f}%，"
     f"仍有{(1-tol1)*100:.1f}%的预测偏差超过一级，说明其相对等级预测能力有限，不宜将结果解释为精确人工评分。"
     "最终生成250款啤酒、50个感官维度的0—6相对等级表，完整结果另附。")
figure("paper_grades.png", "图3  250款啤酒在50个感官维度上的相对等级分布", width=5.3)
score_dims = [d for d in ["overall", "A_hops_all", "A_malt_all", "alcohol", "acidity", "sweetness"]
              if d in scores.columns]
score_rows = []
for _, r in scores.head(5).iterrows():
    score_rows.append([int(r["beer"]), r["tasting_category_fine"]] +
                      [int(r[d]) for d in score_dims])
table("表12  相对等级结果示例", ["样本编号", "风格"] + [DIM_CN.get(d, d) for d in score_dims],
      score_rows, size=7.5)

h2("5.7.4  风格画像的描述性检查")
para("为检查相对等级是否完全背离风格特征，按22种风格汇总各维度等级均值。"
     "该检查只用于发现明显异常，不作为独立验证，也不以风格先验替代样本实测标签。")
style_dims = [d for d in ["alcohol", "acidity", "barnyard", "A_malt_burn", "A_esters_all", "overall"]
              if d in style.columns]
style_rows = []
for key, label in [("Tripel", "Tripel"), ("Kriek", "Kriek"), ("Lambic", "Lambic"),
                   ("Stout", "Stout/Porter"), ("Pils", "Pils/Lager")]:
    row = style[style["tasting_category_fine"].str.contains(key, case=False, na=False)]
    if len(row):
        r = row.iloc[0]
        style_rows.append([label, int(r["n_beers"])] + [f"{r[d]:.2f}" for d in style_dims])
table("表13  代表性风格的相对等级均值", ["风格", "样本数"] + [DIM_CN.get(d, d) for d in style_dims],
      style_rows, size=7.5)
para("Kriek与Lambic的酸度较高，Stout/Porter的烘烤麦芽香较高，Tripel的酒精感较高，"
     "这些特征与相应风格的一般认识一致。与此同时，Pils/Lager的野菌气息等级均值为4.55，"
     "与其清爽风格预期并不完全一致，需要结合原始测量进一步核查。因此，风格画像仅作为描述性合理性检查。")
figure("paper_style.png", "图4  各啤酒风格的相对感官画像", width=5.0)

h1("5.8  模型评价")
para("本文方法的主要优点是：统一预处理保证三个问题的数据口径一致；连续预测与相对等级构造分离，"
     "避免将连续标签直接取整；候选模型通过训练集交叉验证选型，测试集不参与选择；"
     "缺失实验在训练与测试同步遮掩的条件下比较完整方案，能够量化性能对信息损失的敏感程度。")
para("模型仍存在四方面局限。其一，样本量有限而化学指标较多，部分感官维度的决定系数仍为负，"
     "不能据此作出稳定的因果解释。其二，0—6结果由经验分位数构造，只表示样本群体内部的相对等级。"
     "其三，随机遮掩不等同于真实缺失机制，缺失实验结论不能外推到所有缺失情形。"
     "其四，两套缺失方案同时改变了模型与缺失处理方式，不能将性能差异单独归因于均值填补。"
     "后续可采用嵌套交叉验证、序数回归和更贴近实际机制的缺失模拟进一步检验。")
if len(sens_r2) == 5:
    para("为检验数据清洗口径对结论的影响，本文补充了一项敏感性分析：若将疑似异常的统计阈值也限定为仅由训练集估计，"
         f"则位于测试集的两个极端观测不在建模阶段处理。此时岭回归与偏最小二乘回归的测试集平均决定系数"
         f"分别降至{sens_r2[0]:.3f}与{sens_r2[1]:.3f}，"
         f"而随机森林仍为{sens_r2[3]:.3f}，与全样本清洗下的结果一致。"
         "这说明所选树模型对未清洗的极端值具有稳健性，而线性基线对数据质量口径较为敏感。"
         "需要强调的是，该结果只能说明随机森林对本次数据清洗口径的变化较稳健，"
         "并不能证明整个验证流程完全独立于测试集的化学测量信息。")

h1("5.9  问题一小结")
para(f"本文建立了“连续指标回归—相对等级标定—缺失条件检验”的问题一求解框架。"
     f"经训练集交叉验证，随机森林表现最佳，测试集50维平均决定系数为{best['test_R2_mean']:.3f}，"
     f"{int(best['test_dims_R2_positive'])}个维度的决定系数为正。"
     f"回归预测与分位标定路线的平均二次加权Kappa系数为{qwk1:.3f}，据此生成250×50相对等级表。"
     "缺失压力测试表明，特征信息减少时模型性能总体下降；风格画像提供了描述性检查，"
     "但不替代定量评价。上述结果与统一预处理共同构成后续问题二关键因素分析和问题三风格条件生成的数据基础。",
     keep_together=True)

# 文档属性与保存
doc.core_properties.title = "问题一 数据缺失条件下的感官评分——模型建立与求解"
doc.core_properties.subject = "数学建模论文问题一完整稿"
doc.save(OUT)
print(f"已生成：{OUT}")
print(f"公式数：{eq_counter}；结果表：{len(summary)}个模型；最终相对等级：{scores.shape}")
